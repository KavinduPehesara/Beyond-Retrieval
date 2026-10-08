"""Bounded retrieval of Europe PMC's public figure/supplement archive.

Never follows publisher URLs from a paper, extracts archive paths, executes
files, or fetches a URL supplied by the caller. Assets stay outside Git.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import time
import zipfile
from pathlib import Path, PurePosixPath
import xml.etree.ElementTree as ET

import httpx

ASSETS_DIR = Path("data/fulltext_assets")
MAX_ARCHIVE_BYTES = 30_000_000
MAX_FILE_BYTES = 8_000_000
MAX_EXPANDED_BYTES = 60_000_000
MAX_FILES = 150
MAX_TEXT_CHARS = 30_000
RASTER_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".gif": "image/gif", ".webp": "image/webp"}


def valid_pmcid(pmcid: str) -> bool:
    return bool(re.fullmatch(r"PMC\d+", pmcid))


def _checked_zip(data: bytes) -> zipfile.ZipFile:
    archive = zipfile.ZipFile(io.BytesIO(data))
    files = archive.infolist()
    if len(files) > MAX_FILES or sum(f.file_size for f in files) > MAX_EXPANDED_BYTES:
        archive.close()
        raise ValueError("archive exceeds file-count or expanded-size limit")
    return archive


def parse_attachment(data: bytes, suffix: str) -> dict:
    """Preview supported formats. Values are source data, never model claims."""
    result = {"status": "unsupported", "text": "", "tables": [], "note": "Download and inspect this file manually."}
    suffix = suffix.lower()
    try:
        if suffix in (".txt", ".csv", ".tsv", ".xml", ".json"):
            text = data.decode("utf-8-sig", errors="replace")
            result.update(status="extracted", text=text[:MAX_TEXT_CHARS], note="")
            if suffix in (".csv", ".tsv"):
                reader = csv.reader(io.StringIO(text), delimiter="\t" if suffix == ".tsv" else ",")
                rows = []
                for index, row in enumerate(reader):
                    if index >= 100:
                        result["note"] = "Table preview limited to 100 rows. Download the original for all rows."
                        break
                    rows.append([cell[:2000] for cell in row[:100]])
                result["tables"] = [{"label": "Sheet", "rows": rows}]
            if len(text) > MAX_TEXT_CHARS:
                result["note"] += " Text preview truncated at 30,000 characters."
        elif suffix == ".docx":
            with _checked_zip(data) as archive:
                document = ET.fromstring(archive.read("word/document.xml"))
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            paragraphs = ["".join(p.itertext()) for p in document.findall(".//w:p", ns)]
            text = "\n".join(paragraphs)
            result.update(status="extracted", text=text[:MAX_TEXT_CHARS], note="Text preview; formatting is not reproduced.")
            result["tables"] = [
                {"label": f"Table {index + 1}", "rows": [
                    ["".join(cell.itertext())[:2000] for cell in row.findall("w:tc", ns)[:100]]
                    for row in table.findall("w:tr", ns)[:100]]}
                for index, table in enumerate(document.findall(".//w:tbl", ns)[:10])
            ]
            result["note"] += " Table previews limited to 10 tables, 100 rows and 100 cells per row."
            if len(text) > MAX_TEXT_CHARS:
                result["note"] += " Text preview truncated."
        elif suffix == ".xlsx":
            # Check the nested archive before handing it to the workbook reader.
            with _checked_zip(data):
                pass
            import openpyxl
            workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
            try:
                tables = []
                for sheet in workbook.worksheets[:10]:
                    rows = [[str(v)[:2000] if v is not None else "" for v in row]
                            for row in sheet.iter_rows(max_row=min(sheet.max_row or 100, 100),
                                                       max_col=min(sheet.max_column or 100, 100), values_only=True)]
                    tables.append({"label": sheet.title, "rows": rows})
                result.update(status="extracted", tables=tables,
                              note="Preview limited to 10 sheets, 100 rows and 100 columns each. Formula cells use saved values; formulas are not evaluated.")
            finally:
                workbook.close()
        elif suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                result.update(status="needs_human_review", note="Encrypted PDF; inspect the original.")
            else:
                chunks = []
                count = 0
                for page in reader.pages[:30]:
                    text = page.extract_text() or ""
                    chunks.append(text[:MAX_TEXT_CHARS - count])
                    count += len(chunks[-1])
                    if count >= MAX_TEXT_CHARS:
                        break
                text = "\n".join(chunks)
                result.update(status="extracted" if text.strip() else "needs_human_review", text=text,
                              note="PDF text preview: at most 30 pages and 30,000 characters; scanned pages need OCR. Layout and tables require checking.")
    except ImportError as exc:
        result.update(status="unsupported", note=f"Parser dependency unavailable: {exc.name}.")
    except Exception as exc:
        result.update(status="needs_human_review", note=f"Could not parse this file: {type(exc).__name__}.")
    return result


def fetch_assets(pmcid: str, *, client: httpx.Client | None = None, directory: Path | None = None) -> dict:
    if not valid_pmcid(pmcid):
        raise ValueError("invalid PMC identifier")
    directory = (directory or ASSETS_DIR) / pmcid
    manifest = directory / "manifest.json"
    if manifest.exists() and time.time() - manifest.stat().st_mtime < 86400:
        try:
            cached = json.loads(manifest.read_text(encoding="utf-8"))
            if cached.get("version") == 1:
                return cached
        except (OSError, ValueError):
            pass  # An interrupted cache write must not break a new request.
    own = client is None
    client = client or httpx.Client()
    try:
        url = f"https://www.ebi.ac.uk/europepmc/webservices/rest/{pmcid}/supplementaryFiles"
        with client.stream("GET", url, timeout=45.0) as response:
            response.raise_for_status()
            content = bytearray()
            for chunk in response.iter_bytes():
                content.extend(chunk)
                if len(content) > MAX_ARCHIVE_BYTES:
                    raise ValueError("asset archive exceeds 30 MB download limit")
        with _checked_zip(bytes(content)) as archive:
            assets = []
            directory.mkdir(parents=True, exist_ok=True)
            for item in archive.infolist():
                if item.is_dir():
                    continue
                # Original paths are metadata only. Saved filenames are hashes.
                name = PurePosixPath(item.filename.replace("\\", "/")).name
                suffix = Path(name).suffix.lower()
                if item.file_size > MAX_FILE_BYTES:
                    assets.append({"name": name, "status": "source_unavailable", "note": "File exceeds 8 MB limit."})
                    continue
                data = archive.read(item)
                if not data:
                    assets.append({"name": name, "status": "source_unavailable", "note": "Europe PMC returned an empty file."})
                    continue
                asset_id = hashlib.sha256(data).hexdigest()
                (directory / asset_id).write_bytes(data)
                asset = {"name": name, "asset_id": asset_id, "bytes": len(data),
                         "download_url": f"/fulltext/assets/{pmcid}/{asset_id}",
                         "media_type": RASTER_TYPES.get(suffix, "application/octet-stream")}
                if suffix in RASTER_TYPES:
                    asset.update(status="needs_human_review", note="Original image; pixels have not been interpreted.", text="", tables=[])
                else:
                    asset.update(parse_attachment(data, suffix))
                assets.append(asset)
        result = {"version": 1, "status": "extracted" if assets else "not_found", "note": "", "files": assets}
        manifest.write_text(json.dumps(result, ensure_ascii=False), encoding="utf-8")
        return result
    except (httpx.HTTPError, ValueError, zipfile.BadZipFile, RuntimeError) as exc:
        return {"status": "source_unavailable", "note": f"Asset archive unavailable: {exc}", "files": []}
    finally:
        if own:
            client.close()


def find_asset(reference: str, assets: list[dict]) -> dict | None:
    """Match JATS references to archive basenames; ambiguous matches are withheld."""
    basename = PurePosixPath(reference.replace("\\", "/")).name
    exact = [a for a in assets if a["name"] == basename]
    if len(exact) == 1:
        return exact[0]
    stem = Path(basename).stem
    matches = [a for a in assets if Path(a["name"]).stem == stem]
    return matches[0] if len(matches) == 1 else None
