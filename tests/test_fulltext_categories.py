"""Source provenance, missing-source states, assets and attachment boundaries."""
import io
import json
import zipfile

import httpx
import pytest

from slr.adapters.fulltext import parse_jats
from slr.adapters import fulltext_assets as assets
from slr.services.fulltext_categories import enrich_fulltext, unavailable_categories

XML = '''<article xmlns:xlink="http://www.w3.org/1999/xlink" xmlns:m="http://www.w3.org/1998/Math/MathML">
<front><article-meta><article-id pub-id-type="pmc">PMC123</article-id></article-meta></front>
<body><sec><title>Methods</title><p>We used a regression model adjusted for age and sex.</p></sec>
<sec><title>Results</title><p>The mean difference was 4.5 (95% CI 2.1–6.9; P &lt; 0.01).</p>
<table-wrap><label>Table 1</label><caption>Results in mg</caption>
<table><tr><th colspan="2">Treatment</th></tr><tr><td rowspan="2">A</td><td>4.5</td></tr><tr><td>6.9</td></tr></table>
<table-wrap-foot><p>CI means confidence interval.</p></table-wrap-foot></table-wrap>
<fig><label>Figure 1</label><caption>Correlation heat map.</caption><graphic xlink:href="heat.jpg"/></fig>
<fig><label>Figure 2</label><caption>Survival curves.</caption><graphic xlink:href="curve"/></fig>
<disp-formula><label>(1)</label><alternatives><tex-math>y = ax</tex-math><m:math><m:mi>y</m:mi></m:math></alternatives></disp-formula>
</sec><sec><title>Limitations</title><p>We did not measure diet in this trial.</p><p>Future studies should recruit more participants.</p></sec></body>
<back><supplementary-material xlink:href="supp.csv"><label>Supplement 1</label><caption>Extra data</caption></supplementary-material></back>
</article>'''


def zip_bytes(files):
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    return out.getvalue()


def test_rich_source_retains_spans_footnotes_and_one_equation_representation():
    paper = parse_jats(XML)
    assert paper.tables[0].footnotes == ["CI means confidence interval."]
    assert paper.tables[0].cell_spans[0][0]["colspan"] == "2"
    assert paper.equations == ["y = ax"]
    assert "math" in paper.equation_details[0]["mathml"]
    assert paper.supplements[0]["asset_ref"] == "supp.csv"
    assert paper.figures[0].kind == "heat_map"
    assert paper.figures[0].asset_refs == ["heat.jpg"]


def test_all_eight_categories_keep_evidence_verbatim_and_separate_from_tables():
    paper = parse_jats(XML)
    result = enrich_fulltext(paper)
    assert len(result["coverage"]) == 8
    assert len(result["evidence"]["limitations"]) == 2
    for items in result["evidence"].values():
        for item in items:
            assert item["quote"] in paper.body
    assert result["evidence"]["methods"][0]["source"] == "Methods"
    assert "Treatment" not in paper.body
    assert {c["key"]: c["count"] for c in result["coverage"]}["heat_maps"] == 1


def test_unavailable_source_is_never_reported_as_absent_content():
    assert all(c["status"] == "source_unavailable" for c in unavailable_categories("paywalled"))
    fragment = enrich_fulltext(parse_jats("<article><body><p>Only a stub.</p></body></article>"))
    assert fragment["coverage"][0]["status"] == "source_unavailable"


def test_images_are_matched_without_claiming_pixel_analysis_and_supplements_are_previewed():
    files = [
        {"name": "heat.jpg", "asset_id": "a", "media_type": "image/jpeg", "status": "needs_human_review"},
        {"name": "curve.png", "asset_id": "b", "media_type": "image/png", "status": "needs_human_review"},
        {"name": "supp.csv", "asset_id": "c", "status": "extracted", "text": "x,y", "tables": []},
    ]
    result = enrich_fulltext(parse_jats(XML), include_assets=True,
                             asset_loader=lambda _: {"files": files, "note": "", "status": "extracted"})
    assert result["figures"][1]["images"][0]["name"] == "curve.png"
    assert result["figures"][0]["analysis_status"] == "needs_human_review"
    assert result["supplements"][0]["status"] == "extracted"
    assert result["supplements"][0]["text"] == "x,y"


def test_ambiguous_asset_matches_are_not_guessed():
    assert assets.find_asset("curve", [{"name": "curve.png", "asset_id": "a"}, {"name": "curve.jpg", "asset_id": "b"}]) is None


def test_archive_retrieval_uses_fixed_origin_hash_paths_and_keeps_missing_files_explicit(tmp_path):
    archive = zip_bytes({"../../escape.csv": "x,y\n1,2", "empty.pdf": b""})
    requests = []
    def handler(request):
        requests.append(str(request.url))
        return httpx.Response(200, content=archive)
    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        result = assets.fetch_assets("PMC123", client=client, directory=tmp_path)
        assert result["files"][0]["tables"][0]["rows"] == [["x", "y"], ["1", "2"]]
        assert result["files"][1]["status"] == "source_unavailable"
        assert not (tmp_path / "escape.csv").exists()
        assert (tmp_path / "PMC123" / result["files"][0]["asset_id"]).exists()
        assert assets.fetch_assets("PMC123", client=client, directory=tmp_path) == result
    assert len(requests) == 1
    assert requests[0] == "https://www.ebi.ac.uk/europepmc/webservices/rest/PMC123/supplementaryFiles"


def test_asset_size_limits_and_invalid_identifier(tmp_path, monkeypatch):
    with pytest.raises(ValueError):
        assets.fetch_assets("../other", directory=tmp_path)
    monkeypatch.setattr(assets, "MAX_EXPANDED_BYTES", 2)
    with httpx.Client(transport=httpx.MockTransport(lambda _: httpx.Response(200, content=zip_bytes({"big.txt": "abcdef"})))) as client:
        result = assets.fetch_assets("PMC123", client=client, directory=tmp_path)
    assert result["status"] == "source_unavailable"
    assert not (tmp_path / "PMC123" / "manifest.json").exists()


def test_csv_docx_and_spreadsheet_previews_preserve_source_values():
    assert assets.parse_attachment(b"x\ty\n1\t2", ".tsv")["tables"][0]["rows"][1] == ["1", "2"]
    docx = zip_bytes({"word/document.xml": '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Supplement text.</w:t></w:r></w:p></w:body></w:document>'})
    assert "Supplement text." in assets.parse_attachment(docx, ".docx")["text"]
    import openpyxl
    workbook = openpyxl.Workbook()
    workbook.active.append(["Group", "n"])
    workbook.active.append(["A", 12])
    data = io.BytesIO()
    workbook.save(data)
    assert assets.parse_attachment(data.getvalue(), ".xlsx")["tables"][0]["rows"][1] == ["A", "12"]
    assert assets.parse_attachment(b"not a workbook", ".xls")["status"] == "unsupported"


def test_pdf_without_text_requires_review():
    from pypdf import PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    data = io.BytesIO()
    writer.write(data)
    assert assets.parse_attachment(data.getvalue(), ".pdf")["status"] == "needs_human_review"


def test_all_missing_supplements_are_unavailable_even_when_archive_exists():
    result = enrich_fulltext(parse_jats(XML), include_assets=True,
                             asset_loader=lambda _: {"files": [{"name": "supp.csv", "status": "source_unavailable", "note": "empty"}], "note": "", "status": "extracted"})
    category = next(c for c in result["coverage"] if c["key"] == "supplementary_files")
    assert category["status"] == "source_unavailable"


def test_nested_supplement_link_retains_its_parent_label():
    paper = parse_jats('<article xmlns:xlink="http://www.w3.org/1999/xlink"><back><supplementary-material><label>Protocol</label><media xlink:href="protocol.pdf"/></supplementary-material></back></article>')
    assert paper.supplements == [{"label": "Protocol", "caption": "", "asset_ref": "protocol.pdf"}]
